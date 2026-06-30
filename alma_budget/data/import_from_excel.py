#!/usr/bin/env python3
"""Script para importar datos del Excel al módulo alma_budget via XML-RPC.

Uso:
    python3 import_from_excel.py \
        --url https://tu-odoo.com \
        --db nombre_bd \
        --user admin \
        --password tu_password \
        --excel /ruta/al/presupuesto.xlsm

Requiere: openpyxl, xmlrpc (stdlib)
"""

import argparse
import xmlrpc.client
import openpyxl
import sys


def connect(url, db, user, password):
    common = xmlrpc.client.ServerProxy(f'{url}/xmlrpc/2/common')
    uid = common.authenticate(db, user, password, {})
    if not uid:
        raise SystemExit('Autenticación fallida')
    models = xmlrpc.client.ServerProxy(f'{url}/xmlrpc/2/object')
    return uid, models, db


def create_or_update(models, uid, db, password, model, domain, vals):
    existing = models.execute_kw(db, uid, password, model, 'search', [domain])
    if existing:
        models.execute_kw(db, uid, password, model, 'write', [existing, vals])
        return existing[0]
    return models.execute_kw(db, uid, password, model, 'create', [vals])


def import_articulos(wb, models, uid, db, password):
    """Importa tejidos desde pestaña ARTICULOS."""
    ws = wb['ARTICULOS']
    print('Importando tejidos...')
    count = 0
    for row in range(10, ws.max_row + 1):
        ref1 = ws.cell(row, 4).value
        ref2 = ws.cell(row, 5).value
        name = ws.cell(row, 6).value
        shrink_len = ws.cell(row, 7).value or 0.0
        shrink_wid = ws.cell(row, 8).value or 0.0

        if not (ref1 and ref2 and name):
            continue

        vals = {
            'ref1': int(ref1),
            'ref2': str(ref2).strip(),
            'name': str(name).strip(),
            'shrinkage_length': float(shrink_len),
            'shrinkage_width': float(shrink_wid),
        }
        create_or_update(
            models, uid, db, password, 'alma.fabric',
            [('ref1', '=', vals['ref1']), ('ref2', '=', vals['ref2'])],
            vals
        )
        count += 1

    print(f'  → {count} tejidos importados.')


def import_colores(wb, models, uid, db, password):
    """Importa colores desde pestaña COLORES."""
    ws = wb['COLORES']
    print('Importando colores...')
    count = 0

    # Obtener mapa ref_combined → id de alma.fabric
    fabrics = models.execute_kw(
        db, uid, password, 'alma.fabric', 'search_read',
        [[]], {'fields': ['id', 'ref_combined']}
    )
    fabric_map = {f['ref_combined']: f['id'] for f in fabrics}

    for row in range(7, ws.max_row + 1):
        ref2_combined = ws.cell(row, 3).value  # col C = ej. "1TVR"
        color_code = ws.cell(row, 6).value
        color_name = ws.cell(row, 7).value
        cost = ws.cell(row, 8).value
        width = ws.cell(row, 9).value
        composition = ws.cell(row, 10).value

        if not (ref2_combined and color_code and cost):
            continue

        ref2_combined = str(ref2_combined).strip()
        fabric_id = fabric_map.get(ref2_combined)
        if not fabric_id:
            print(f'  ! Tejido no encontrado para REF={ref2_combined}, saltando fila {row}')
            continue

        vals = {
            'fabric_id': fabric_id,
            'color_code': int(color_code),
            'color_name': str(color_name).strip() if color_name else '',
            'cost': float(cost),
            'width': float(width) if width else 0.0,
            'composition': str(composition).strip() if composition else '',
        }
        create_or_update(
            models, uid, db, password, 'alma.fabric.color',
            [('fabric_id', '=', fabric_id), ('color_code', '=', int(color_code))],
            vals
        )
        count += 1

    print(f'  → {count} colores importados.')


def import_empalmes(wb, models, uid, db, password):
    """Importa tabla de empalmes desde pestaña EMPALME."""
    ws = wb['EMPALME']
    print('Importando empalmes...')
    count = 0
    for row in range(2, ws.max_row + 1):
        fw = ws.cell(row, 1).value   # Ancho tejido
        mw = ws.cell(row, 2).value   # Ancho mantel
        emp_len = ws.cell(row, 4).value  # Largo empalme

        if not (fw and mw and emp_len):
            continue

        vals = {
            'fabric_width_cm': int(fw),
            'mantel_width_cm': int(mw),
            'empalme_length_cm': int(emp_len),
        }
        create_or_update(
            models, uid, db, password, 'alma.empalme',
            [('fabric_width_cm', '=', int(fw)), ('mantel_width_cm', '=', int(mw))],
            vals
        )
        count += 1

    print(f'  → {count} empalmes importados.')


def main():
    parser = argparse.ArgumentParser(description='Importar Excel a Odoo alma_budget')
    parser.add_argument('--url', required=True)
    parser.add_argument('--db', required=True)
    parser.add_argument('--user', default='admin')
    parser.add_argument('--password', required=True)
    parser.add_argument('--excel', required=True)
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.excel, keep_vba=True, data_only=True)
    uid, models, db = connect(args.url, args.db, args.user, args.password)

    import_articulos(wb, models, uid, db, args.password)
    import_colores(wb, models, uid, db, args.password)
    import_empalmes(wb, models, uid, db, args.password)

    print('\n✅ Importación completada.')


if __name__ == '__main__':
    main()
